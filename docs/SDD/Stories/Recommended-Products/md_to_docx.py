import sys, os, re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def shade(cell, color):
    e = OxmlElement('w:shd'); e.set(qn('w:fill'), color); e.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(e)

def add_runs(paragraph, text):
    # handle **bold**, `code`
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for p in parts:
        if not p:
            continue
        if p.startswith('**') and p.endswith('**'):
            r = paragraph.add_run(p[2:-2]); r.bold = True
        elif p.startswith('`') and p.endswith('`'):
            r = paragraph.add_run(p[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
        else:
            paragraph.add_run(p)

def convert(md_path, docx_path, title):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)

    # title
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run(title); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(44,62,80)

    lines = open(md_path, encoding='utf-8').read().split('\n')
    i = 0
    in_code = False
    code_buf = []
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # code fence
        if stripped.startswith('```'):
            if not in_code:
                in_code = True; code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                run = p.add_run('\n'.join(code_buf))
                run.font.name = 'Consolas'; run.font.size = Pt(9)
                shade_para(p, 'F2F3F4')
            i += 1; continue
        if in_code:
            code_buf.append(line); i += 1; continue

        # table
        if stripped.startswith('|') and i+1 < len(lines) and re.match(r'^\|[\s:\-|]+\|?$', lines[i+1].strip()):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                tbl_lines.append(lines[i].strip()); i += 1
            rows = []
            for tl in tbl_lines:
                if re.match(r'^\|[\s:\-|]+\|?$', tl):
                    continue
                cells = [c.strip() for c in tl.strip('|').split('|')]
                rows.append(cells)
            if rows:
                ncol = len(rows[0])
                table = doc.add_table(rows=len(rows), cols=ncol); table.style = 'Table Grid'
                table.alignment = WD_TABLE_ALIGNMENT.LEFT
                for ri, rdata in enumerate(rows):
                    for ci in range(ncol):
                        cell = table.rows[ri].cells[ci]
                        cell.text = ''
                        para = cell.paragraphs[0]
                        add_runs(para, rdata[ci] if ci < len(rdata) else '')
                        for rr in para.runs:
                            rr.font.size = Pt(9)
                            if ri == 0: rr.bold = True; rr.font.color.rgb = RGBColor(255,255,255)
                        if ri == 0: shade(cell, '2C3E50')
                        elif ri % 2 == 0: shade(cell, 'F4F6F7')
            doc.add_paragraph()
            continue

        # headings
        if stripped.startswith('#'):
            m = re.match(r'^(#+)\s*(.*)$', stripped)
            level = min(len(m.group(1)), 4)
            txt = m.group(2).strip()
            if level == 1:
                # skip duplicate top title (already added) but keep as H1 if different
                if txt.lower().replace('—','-') not in title.lower().replace('—','-'):
                    doc.add_heading(txt, level=1)
            else:
                doc.add_heading(txt, level=level-0 if level>1 else 1)
            i += 1; continue

        # horizontal rule
        if stripped in ('---', '***', '___'):
            i += 1; continue

        # blockquote
        if stripped.startswith('>'):
            p = doc.add_paragraph()
            run = p.add_run('Note: '); run.bold = True; run.font.color.rgb = RGBColor(41,128,185)
            add_runs(p, stripped.lstrip('> ').replace('Note: ','').replace('Implementation note: ',''))
            i += 1; continue

        # checkbox / bullet
        m_cb = re.match(r'^- \[ \]\s*(.*)$', stripped)
        if m_cb:
            p = doc.add_paragraph(style='List Bullet'); add_runs(p, m_cb.group(1)); i += 1; continue
        m_b = re.match(r'^[-*]\s+(.*)$', stripped)
        if m_b:
            p = doc.add_paragraph(style='List Bullet'); add_runs(p, m_b.group(1)); i += 1; continue
        m_n = re.match(r'^\d+\.\s+(.*)$', stripped)
        if m_n:
            p = doc.add_paragraph(style='List Number'); add_runs(p, m_n.group(1)); i += 1; continue

        # blank
        if not stripped:
            i += 1; continue

        # paragraph
        p = doc.add_paragraph(); add_runs(p, stripped)
        i += 1

    doc.save(docx_path)
    print('Saved:', docx_path)

def shade_para(paragraph, color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),color)
    pPr.append(shd)

if __name__ == '__main__':
    here = os.path.dirname(os.path.abspath(__file__))
    jobs = [
        ('EC-Recommended-Products-GraphQL-Story.md', 'EC-Recommended-Products-GraphQL-Story.docx',
         'ACCS Story — GraphQL API for Recommended Products'),
        ('EC-Recommended-Products-Block-Frontend-Story.md', 'EC-Recommended-Products-Block-Frontend-Story.docx',
         'EDS Front-End Story — Recommended Products Block'),
    ]
    for md, dx, title in jobs:
        convert(os.path.join(here, md), os.path.join(here, dx), title)
