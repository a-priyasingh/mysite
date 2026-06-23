import csv, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def shade(cell, color):
    e = OxmlElement('w:shd'); e.set(qn('w:fill'), color); e.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(e)

def convert(csv_path, docx_path, title):
    rows = list(csv.reader(open(csv_path, encoding='utf-8')))
    headers = rows[0]
    data = rows[1:]

    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)

    t = doc.add_paragraph()
    r = t.add_run(title); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(44,62,80)
    sub = doc.add_paragraph()
    rs = sub.add_run(f'{len(data)} blocks. One section per block; fields shown as labeled rows.')
    rs.italic = True; rs.font.size = Pt(9); rs.font.color.rgb = RGBColor(86,101,115)

    # Block name is column 0; render each row as its own heading + 2-col detail table
    for row in data:
        block_name = row[0] if row else ''
        doc.add_heading(block_name, level=1)
        # 2-col table: field label | value, for columns 1..n
        pairs = [(headers[i], row[i] if i < len(row) else '') for i in range(1, len(headers))]
        tbl = doc.add_table(rows=len(pairs), cols=2); tbl.style = 'Table Grid'
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        for ri, (label, val) in enumerate(pairs):
            lc = tbl.rows[ri].cells[0]; vc = tbl.rows[ri].cells[1]
            lc.text = label
            for p in lc.paragraphs:
                for rr in p.runs:
                    rr.bold = True; rr.font.size = Pt(9); rr.font.color.rgb = RGBColor(255,255,255)
            shade(lc, '2C3E50')
            vc.text = val
            for p in vc.paragraphs:
                for rr in p.runs: rr.font.size = Pt(9)
            shade(vc, 'F4F6F7' if ri % 2 == 0 else 'FFFFFF')
        # column widths
        from docx.shared import Cm
        for cells in tbl.columns[0].cells: cells.width = Cm(4.5)
        for cells in tbl.columns[1].cells: cells.width = Cm(12.5)
        doc.add_paragraph()

    doc.save(docx_path)
    print('Saved:', docx_path, '|', len(data), 'blocks')

if __name__ == '__main__':
    here = os.path.dirname(os.path.abspath(__file__))
    jobs = [
        ('GSUSA-Home-Page-Block-Stories.csv', 'GSUSA-Home-Page-Block-Stories.docx',
         'GSUSA Home Page Block Stories — Desktop'),
        ('GSUSA-Home-Page-Block-Stories-Mobile.csv', 'GSUSA-Home-Page-Block-Stories-Mobile.docx',
         'GSUSA Home Page Block Stories — Mobile'),
    ]
    for c, d, ti in jobs:
        convert(os.path.join(here, c), os.path.join(here, d), ti)
