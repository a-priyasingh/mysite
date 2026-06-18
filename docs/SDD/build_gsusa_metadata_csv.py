import csv, os

HEADERS = [
 "Custom or Out of the Box (OOTB)","Section Title","Metadata Field Label","Technical Property Name",
 "Summary","Tool Tips","GSUSA Feedback (Keep or Remove)","Feedback Resolved","Required (Y/N)",
 "Enabled/Disabled","Extract/Autopopulate OR Manual","Type","Multi / Single Select",
 "Values (Dropdown Items or Tag Path)","Metadata Schema (*Global = Default, all assets)","Folders Applied",
 "Original Source (WebDAM/SharePoint/NetSuite)","girlscoutshop.com","girlscouts.org / Council Sites",
 "Email / Social","Previous Mapping (WebDAM Field)","Consolidation Notes",
]

# Each row: OOTB/Custom, Section, Label, Property, Summary, ToolTip, Feedback, Resolved, Req,
# Enabled, Extract/Manual, Type, Multi/Single, Values, Schema, Folders, Source, shop, org, email, prevMap, notes
R = []
def row(*a): R.append(list(a) + [""]*(len(HEADERS)-len(a)))

# ---------------- BASIC TAB ----------------
row("OOTB AEM Field","Metadata (Basic)","Title","./jcr:content/metadata/dc:title","Primary name representing the asset.","Enter a concise descriptive title.","","","Y","Enabled","Manual","Open Text","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","x","x","x","Headline (Caption/Title)","Maps from WebDAM Headline; NetSuite ALT text can seed this.")
row("OOTB AEM Field","Metadata (Basic)","Description","./jcr:content/metadata/dc:description","Long description / abstract of the asset.","Describe the asset content and context.","","","N","Enabled","Manual","Open Text","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","x","x","x","Caption/Abstract (Long Description)","")
row("Custom AEM Field","Metadata (Basic)","Asset Type","./jcr:content/metadata/gs:assetType","Classifies the kind of asset.","Select the asset type.","","","Y","Enabled","Manual","Dropdown","Single","Product Image; Lifestyle Photo; Logo; Icon/Illustration; Badge/Patch; Document/PDF; Video; Banner/Creative; Infographic; Brand Guideline","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","x","x","x","Image Type (Asset Type)","PROPOSED starter values - confirm with GSUSA (Q8).")
row("OOTB AEM Field","Metadata (Basic)","Tags / Keywords","./jcr:content/metadata/cq:tags","Controlled taxonomy tags for classification.","Apply tags from the girlscouts taxonomy.","","","N","Enabled","Manual","Tag List","Multiselect","/content/cq:tags/girlscouts","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","x","x","x","Keywords","")
row("OOTB AEM Field","Metadata (Basic)","Smart Tags","./jcr:content/metadata/dam:smartTags","AI-generated descriptive tags.","Auto-applied by Asset Compute.","","","N","Enabled","Autopopulate (AI)","Tag List","Multiselect","Smart Tags","Global Metadata Schema (All Assets)","/content/dam/gsusa","","","","","","")
row("OOTB AEM Field","Metadata (Basic)","Language","./jcr:content/metadata/dc:language","Language of the asset content.","Select the asset language.","","","N","Enabled","Manual","Dropdown","Multiselect","English; Spanish","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","x","x","x","Language Identifier","Localization handled via metadata, not folders.")
row("OOTB AEM Field","Metadata (Basic)","Created Date","./jcr:content/jcr:created","Date the asset was added to AEM.","System-populated.","","","N","Enabled","Autopopulate (System)","Date","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa","","","","","Date Created","")

# ---------------- RIGHTS & CONDITIONS ----------------
row("Custom AEM Field","Rights & Conditions","Owner of Asset","./jcr:content/metadata/gs:owner","Who owns the asset (GSUSA or external).","Select asset owner.","","","Y","Enabled","Manual","Dropdown","Single","GS; External","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","x","x","x","","")
row("OOTB AEM Field","Rights & Conditions","Usage Rights","./jcr:content/metadata/xmpRights:UsageTerms","Permitted usage / rights for the asset.","Select usage rights.","","","Y","Enabled","Manual","Dropdown","Multiselect","Internal Use Only; Web/Digital; Print; Social Media; Email; All Channels; Council Use; Restricted","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","x","x","x","Rights Usage Terms (Asset Usage)","PROPOSED starter values - confirm with GSUSA (Q8).")
row("Custom AEM Field","Rights & Conditions","License Type","./jcr:content/metadata/gs:rightsLicenseType","Type of license governing the asset.","Select license type.","","","N","Enabled","Manual","Dropdown","Single","Royalty-Free; Rights-Managed; Owned/In-House; Licensed (Time-Limited); Creative Commons; Stock; Model-Released","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","","","","Usage Terms (Usage Rights)","PROPOSED starter values - confirm with GSUSA (Q8).")
row("OOTB AEM Field","Rights & Conditions","Copyright","./jcr:content/metadata/dc:rights","Copyright notice.","Enter copyright text.","","","N","Enabled","Manual","Open Text","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","","","","Copyright Notice","")
row("OOTB AEM Field","Rights & Conditions","Photographer / Creator","./jcr:content/metadata/dc:creator","Person who created the asset.","Enter creator name.","","","N","Enabled","Manual","Open Text","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","","","","By-line (Photographer/Creator)","")
row("OOTB AEM Field","Rights & Conditions","Rights Notes","./jcr:content/metadata/photoshop:Instructions","Special rights instructions.","Enter rights notes.","","","N","Enabled","Manual","Open Text","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","","","","Special Instructions (Rights Notes)","")
row("OOTB AEM Field","Rights & Conditions","Expiration Date","./jcr:content/metadata/prism:expirationDate","Date after which the asset must not be used.","Set expiry / valid-to date.","","","N","Enabled","Manual","Date","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa","","","","","","Drives expiry notifications.")

# ---------------- MARKETING & CAMPAIGN ----------------
row("Custom AEM Field","Marketing & Campaign","Campaign / Project Name","./jcr:content/metadata/gs:campaign","Associated campaign or project.","Apply campaign tag.","","","N","Enabled","Manual","Tag List","Multiselect","/content/cq:tags/girlscouts/campaign","Global Metadata Schema (All Assets)","/content/dam/gsusa/marketing","WebDAM","x","x","x","Campaign/Project Name","Use single parent tag; avoid duplicates (MOM).")
row("Custom AEM Field","Marketing & Campaign","Campaign Type","./jcr:content/metadata/gs:campaignType","Type/category of campaign.","Apply campaign type tag.","","","N","Enabled","Manual","Tag List","Multiselect","/content/cq:tags/girlscouts/campaign-type","Global Metadata Schema (All Assets)","/content/dam/gsusa/marketing","","x","x","x","","")
row("Custom AEM Field","Marketing & Campaign","Campaign Activation Year","./jcr:content/metadata/gs:campaignYear","Year the campaign was activated.","Select activation year.","","","N","Enabled","Manual","Dropdown","Single","2023; 2024; 2025; 2026","Global Metadata Schema (All Assets)","/content/dam/gsusa/marketing","WebDAM","","","","Campaign/Project Activation Year","")
row("Custom AEM Field","Marketing & Campaign","Content Type","./jcr:content/metadata/gs:contentType","Marketing content format.","Select content type.","","","N","Enabled","Manual","Dropdown","Single","Banner; Email; Social; Landing Page","Global Metadata Schema (All Assets)","/content/dam/gsusa/marketing","","x","x","x","","")
row("Custom AEM Field","Marketing & Campaign","Marketing Objective","./jcr:content/metadata/gs:marketingObjective","Goal the asset supports.","Apply marketing objective tag.","","","N","Enabled","Manual","Tag List","Multiselect","Recruitment; Renewal; Retention; Awareness; Fundraising; Volunteer Engagement; Product Sales; Brand Building","Global Metadata Schema (All Assets)","/content/dam/gsusa/marketing","","","","x","","PROPOSED starter taxonomy - confirm with GSUSA (Q9).")
row("Custom AEM Field","Marketing & Campaign","Promotion Type","./jcr:content/metadata/gs:promotionType","Type of promotion.","Apply promotion type tag.","","","N","Enabled","Manual","Tag List","Multiselect","Discount/Sale; BOGO; Free Shipping; Bundle; New Arrival; Clearance; Seasonal Offer; Member Exclusive","Global Metadata Schema (All Assets)","/content/dam/gsusa/marketing","","x","","x","","PROPOSED starter taxonomy - confirm with GSUSA (Q9).")
row("Custom AEM Field","Marketing & Campaign","Seasonal Type","./jcr:content/metadata/gs:seasonalType","Seasonal association.","Apply seasonal tag.","","","N","Enabled","Manual","Tag List","Multiselect","Spring; Summer; Fall; Winter; Back-to-Troop; Holiday; Cookie Season; Bridging","Global Metadata Schema (All Assets)","/content/dam/gsusa/marketing","","x","x","x","","PROPOSED starter taxonomy - confirm with GSUSA (Q9).")
row("Custom AEM Field","Marketing & Campaign","Event Name","./jcr:content/metadata/gs:eventName","Associated event.","Apply event tag.","","","N","Enabled","Manual","Tag List","Multiselect","/content/cq:tags/girlscouts/events","Global Metadata Schema (All Assets)","/content/dam/gsusa/Events","WebDAM","","x","x","Event Name","")

# ---------------- AUDIENCE & PROGRAM ----------------
row("Custom AEM Field","Audience & Program","Audience Role","./jcr:content/metadata/gs:audienceRole","Intended audience.","Select audience role.","","","N","Enabled","Manual","Dropdown","Multiselect","Girl Scouts; Troop Leaders; Parents; Volunteers","Global Metadata Schema (All Assets)","/content/dam/gsusa","","x","x","x","","")
row("Custom AEM Field","Audience & Program","Gender","./jcr:content/metadata/gs:genderAudience","Gender audience for product/apparel.","Select gender.","","","N","Enabled","Manual","Dropdown","Single","Women; Men; Unisex","Global Metadata Schema (All Assets)","/content/dam/gsusa","","x","","","","")
row("Custom AEM Field","Audience & Program","Program Level","./jcr:content/metadata/gs:programLevel","Girl Scout program/grade level.","Apply program level tag.","","","N","Enabled","Manual","Tag List","Multiselect","Daisy; Brownie; Junior; Cadette; Senior; Ambassador","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","x","x","x","","From WebDAM Age Level keywords.")
row("Custom AEM Field","Audience & Program","Pillars (Themes)","./jcr:content/metadata/gs:pillars","Program pillar / theme.","Apply pillar tag.","","","N","Enabled","Manual","Tag List","Multiselect","Entrepreneurship; Life Skills; Outdoors; STEM","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","","x","x","Pillars","")
row("Custom AEM Field","Audience & Program","Badge Type","./jcr:content/metadata/gs:badgeType","Type of badge.","Apply badge type tag.","","","N","Enabled","Manual","Tag List","Multiselect","Earned Badge; Earned Pin; Fun Patch; Journey Award; Skill-Building Badge; Cookie Business Badge","Global Metadata Schema (All Assets)","/content/dam/gsusa","","x","x","","","PROPOSED starter values - confirm with GSUSA.")
row("Custom AEM Field","Audience & Program","Highest Award","./jcr:content/metadata/gs:highestAward","Highest award association.","Apply award tag.","","","N","Enabled","Manual","Tag List","Multiselect","Bronze; Silver; Gold","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","","x","","","")
row("Custom AEM Field","Audience & Program","Activity","./jcr:content/metadata/gs:activity","Associated activity.","Apply activity tag.","","","N","Enabled","Manual","Tag List","Multiselect","/content/cq:tags/girlscouts/activity","Global Metadata Schema (All Assets)","/content/dam/gsusa/activity","","","x","x","","")

# ---------------- PHOTOGRAPHY & LOCATION ----------------
row("Custom AEM Field","Photography & Location","Photo Shoot Code","./jcr:content/metadata/gs:photoShootCode","Identifier for the photo shoot.","Enter shoot code (e.g. SHOOT-2026-13).","","","N","Enabled","Manual","Open Text","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa/photoshoots","SharePoint","","x","x","","Lifestyle photography from SharePoint.")
row("Custom AEM Field","Photography & Location","Mood","./jcr:content/metadata/gs:mood","Mood/tone of the imagery.","Select mood.","","","N","Enabled","Manual","Dropdown","Multiselect","Adventurous; Solemn; Sporty","Global Metadata Schema (All Assets)","/content/dam/gsusa/photoshoots","","","x","x","","")
row("Custom AEM Field","Photography & Location","Location Type","./jcr:content/metadata/gs:locationType","Type of location depicted.","Select location type.","","","N","Enabled","Manual","Dropdown","Single","Camp; Dream Lab; Office","Global Metadata Schema (All Assets)","/content/dam/gsusa/photoshoots","","","x","","","")
row("Custom AEM Field","Photography & Location","Location Name","./jcr:content/metadata/gs:locationName","Specific location name.","Enter location name.","","","N","Enabled","Manual","Open Text","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa/photoshoots","","","","","","Council/camp/dream lab name.")
row("OOTB AEM Field","Photography & Location","City","./jcr:content/metadata/photoshop:City","City where the asset was created.","Enter city.","","","N","Enabled","Manual","Open Text","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa/photoshoots","WebDAM","","","","City","")
row("OOTB AEM Field","Photography & Location","State","./jcr:content/metadata/photoshop:State","State/province.","Select state.","","","N","Enabled","Manual","Dropdown","Single","US States","Global Metadata Schema (All Assets)","/content/dam/gsusa/photoshoots","WebDAM","","","","Province/State","")

# ---------------- COUNCIL ----------------
row("Custom AEM Field","Council","Council Asset?","./jcr:content/metadata/gs:council","Whether the asset is council-owned.","Yes/No.","","","N","Enabled","Manual","Dropdown","Single","Yes; No","Global Metadata Schema (All Assets)","/content/dam/gsusa/councils","","","x","","","Promote to required in Council schema (Phase 2).")
row("Custom AEM Field","Council","Council ID","./jcr:content/metadata/gs:councilIds","Council code(s) associated.","Enter council ID(s).","","","N","Enabled","Manual","Open Text","Multiselect","Council codes (e.g. 603, 607)","Global Metadata Schema (All Assets)","/content/dam/gsusa/councils","","","x","","","")
row("Custom AEM Field","Council","Council Region","./jcr:content/metadata/gs:councilRegion","Region grouping for councils.","Apply council region tag.","","","N","Enabled","Manual","Tag List","Multiselect","/content/cq:tags/girlscouts/council","Global Metadata Schema (All Assets)","/content/dam/gsusa/councils","","","x","","","")

# ---------------- COMMERCE (PRODUCT SCHEMA) ----------------
row("Custom AEM Field","Commerce","Eligible for Commerce","./jcr:content/metadata/commerce:isCommerce","Whether the asset is used on girlscoutshop.com.","Yes/No.","","","Y","Enabled","Manual","Dropdown","Single","Yes; No","Product / Commerce (Light) Schema","/content/dam/gsusa/products","NetSuite/WebDAM","x","","","","Required in Product schema only.")
row("Custom AEM Field","Commerce","Product Category","./jcr:content/metadata/gs:productCategory","Product taxonomy category.","Apply product category tag.","","","N","Enabled","Manual","Tag List","Multiselect","/content/cq:tags/girlscouts/product","Product / Commerce (Light) Schema","/content/dam/gsusa/products","WebDAM","x","","","","Taxonomy pending completion (Q9).")
row("Custom AEM Field","Commerce","Collection","./jcr:content/metadata/gs:collection","Product collection grouping.","Apply collection tag.","","","N","Enabled","Manual","Tag List","Multiselect","/content/cq:tags/girlscouts/collection","Product / Commerce (Light) Schema","/content/dam/gsusa/products","WebDAM","x","","","","")
row("Custom AEM Field","Commerce","Product SKU","./jcr:content/metadata/commerce:skus","Commerce SKU(s) linked to the asset.","Enter SKU number(s).","","","N","Enabled","Extract/Autopopulate","Open Text","Multiselect","","Product / Commerce (Light) Schema","/content/dam/gsusa/products","NetSuite","x","","","Product SKU Number","From NetSuite/commerce.")
row("Custom AEM Field","Commerce","Position","./jcr:content/metadata/commerce:positions","Display position/order for the product.","Enter position number.","","","N","Enabled","Manual","Number","Single","","Product / Commerce (Light) Schema","/content/dam/gsusa/products","","x","","","","")
row("Custom AEM Field","Commerce","Role","./jcr:content/metadata/commerce:roles","Role of the image (e.g. main, alt).","Select role.","","","N","Enabled","Manual","Dropdown","Single","Main/Primary; Alternate; Swatch; Lifestyle; Detail/Zoom; Back; Packaging; Size Chart","Product / Commerce (Light) Schema","/content/dam/gsusa/products","","x","","","","PROPOSED starter values - confirm with GSUSA (Q8).")

# ---------------- MIGRATION / LEGACY ----------------
row("Custom AEM Field","Migration / Legacy","Legacy Import Date","./jcr:content/metadata/gs:legacyImportDate","Date the asset was imported from a legacy system.","System-populated at migration.","","","N","Enabled","Autopopulate (Migration)","Date","Single","","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM/SharePoint/NetSuite","","","","","Admin/hidden field.")
row("Custom AEM Field","Migration / Legacy","Legacy Source System","./jcr:content/metadata/gs:legacySourceSystem","Original source system of the asset.","Select source.","","","N","Enabled","Autopopulate (Migration)","Dropdown","Single","WebDAM; SharePoint; NetSuite","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM/SharePoint/NetSuite","","","","","Admin/hidden field.")
row("Custom AEM Field","Migration / Legacy","Sync to AEM","./jcr:content/metadata/gs:syncToAem","Sync flag from legacy DAM.","","","","N","Disabled","Autopopulate","Dropdown","Single","Yes; No","Global Metadata Schema (All Assets)","/content/dam/gsusa","WebDAM","","","","Sync to AEM","Likely deprecated post-migration.")

here = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(here, 'GSUSA-Asset-Metadata-Schema.csv')
with open(out, 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f); w.writerow(HEADERS)
    for r in R: w.writerow(r)
print('Saved:', out, '| fields:', len(R))

# ---------------- XLSX (formatted) ----------------
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Future-State Asset Metadata'

hdr_fill = PatternFill('solid', fgColor='2C3E50')
hdr_font = Font(bold=True, color='FFFFFF', size=10)
sec_fill = PatternFill('solid', fgColor='D6EAF8')
sec_font = Font(bold=True, size=10, color='1A5276')
thin = Side(style='thin', color='D5D8DC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(vertical='top', wrap_text=True)

# Title row
ws.cell(1, 1, 'GSUSA Asset Metadata Schema - Future State (girlscoutshop.com)')
ws.cell(1, 1).font = Font(bold=True, size=13, color='2C3E50')

# Header row at row 2
for ci, hname in enumerate(HEADERS, start=1):
    c = ws.cell(2, ci, hname)
    c.fill = hdr_fill; c.font = hdr_font; c.alignment = wrap; c.border = border

# Data rows with section grouping
cur_section = None
rownum = 3
for r in R:
    section = r[1]
    if section != cur_section:
        # section banner row spanning all columns
        ws.cell(rownum, 1, section)
        for ci in range(1, len(HEADERS)+1):
            cc = ws.cell(rownum, ci); cc.fill = sec_fill; cc.border = border
            if ci == 1: cc.font = sec_font
        ws.merge_cells(start_row=rownum, start_column=1, end_row=rownum, end_column=len(HEADERS))
        cur_section = section
        rownum += 1
    for ci, val in enumerate(r, start=1):
        cc = ws.cell(rownum, ci, val); cc.alignment = wrap; cc.border = border; cc.font = Font(size=9)
    rownum += 1

# Column widths
widths = [16,18,22,42,40,28,22,14,11,13,22,14,16,46,30,28,26,14,18,14,28,40]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = 'C3'  # freeze header + first two label cols
ws.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}2"

outx = os.path.join(here, 'GSUSA-Asset-Metadata-Schema.xlsx')
wb.save(outx)
print('Saved:', outx)
