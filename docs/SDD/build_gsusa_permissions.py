import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ---- shared styles ----
HDR_FILL = PatternFill('solid', fgColor='2C3E50')
HDR_FONT = Font(bold=True, color='FFFFFF', size=10)
SEC_FILL = PatternFill('solid', fgColor='D6EAF8')
SEC_FONT = Font(bold=True, size=10, color='1A5276')
INT_FILL = PatternFill('solid', fgColor='EAF7EE')   # internal
EXT_FILL = PatternFill('solid', fgColor='FDF2E3')   # external
thin = Side(style='thin', color='D5D8DC')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(vertical='top', wrap_text=True)
CENTER = Alignment(vertical='center', horizontal='center', wrap_text=True)

def style_header(ws, row, ncols):
    for c in range(1, ncols+1):
        cell = ws.cell(row, c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT; cell.alignment = WRAP; cell.border = BORDER

def put_row(ws, rownum, values, fill=None, center_from=None):
    for ci, v in enumerate(values, start=1):
        cell = ws.cell(rownum, ci, v)
        cell.alignment = CENTER if (center_from and ci >= center_from) else WRAP
        cell.border = BORDER; cell.font = Font(size=9)
        if fill: cell.fill = fill

# =====================================================================
# SHEET 1 - User Type Summary (matches GSUSA Permissions format)
# =====================================================================
ws1 = wb.active
ws1.title = 'Permission Matrix - Summary'

ws1.cell(1,1,'GSUSA AEM Assets (DAM) - User Permission Matrix').font = Font(bold=True, size=13, color='2C3E50')
ws1.cell(2,1,'Based on DAM Strategy Workbook + Sprint 0 Refinement MOM (17 Jun 2026). Authentication: Okta (internal girlscouts.org; external partners on separate domains).').font = Font(italic=True, size=9, color='566573')

H1 = ['User Type','Internal / External','Authentication','Access Method','Allowed Activities in Adobe Assets',
      'Folder-Level Permissions','Tag/Metadata Rights','Approx. Users','Notes / Decisions (per MOM)']
hrow = 4
for ci,h in enumerate(H1, start=1): ws1.cell(hrow,ci,h)
style_header(ws1, hrow, len(H1))

rows1 = [
 ('GSUSA Content Creators','Internal','Okta (SSO)','Direct AEM Assets',
  'Create/upload assets, update & delete assets, modify metadata, apply tags. Full superset access.',
  'Full read/write/create/delete on ALL /content/dam/gsusa folders.',
  'Apply existing tags + metadata. Create new tags: NO (admin-only).',
  '~20 (to confirm vs license)','Full superset access to all DAM folders (MOM 2.5).', 'int'),
 ('GSUSA Content Managers','Internal','Okta (SSO)','Direct AEM Assets',
  'Create/update assets & metadata in Marketing; view-only elsewhere; apply tags.',
  'Full read/write on /content/dam/gsusa/marketing. View-only on all other folders.',
  'Apply existing tags. Create new tags: NO.',
  'Subset of internal','Write limited to Marketing; view-only elsewhere; some folders may be hidden (MOM 2.5).', 'int'),
 ('GSUSA Content Users','Internal','Okta (SSO)','Direct AEM Assets / Content Hub',
  'Search, view, and download assets only.',
  'Read + download on all folders (except explicitly restricted GSUSA folders).',
  'None (read-only).',
  'Largest internal group','View/download consumers; no edit rights.', 'int'),
 ('Admin / IT','Internal','Okta (SSO)','Direct AEM Assets (admin)',
  'Manage users, groups, permissions/ACLs, folder structure, schemas, processing profiles, audit.',
  'Full access (incl. Read-ACL / Edit-ACL / Replicate) on all folders.',
  'Create/manage tags, schemas, taxonomy. Tag creation restricted to 1-2 admins (MOM 2.5).',
  '1-2','Only admins create new tags & manage configuration.', 'int'),
 ('Council Collaborators','External','Okta (separate domain)','Content Hub (no direct AEM)',
  'Upload council-specific assets, update metadata, apply tags on own assets, submit for approval; view/download shared GSUSA assets.',
  'Full read/write/edit/delete WITHIN own council folder only. View shared GSUSA folders (global, products, marketing). Cannot view other councils.',
  'Apply/update tags + metadata on own council assets. Create new tags: NO.',
  '>=1, likely 2-3 per council (110 councils)','Okta separate domain. No direct AEM access - upload/view/download via Content Hub (Adobe to confirm CH upload). Council 100 (Girl Scout Central) = GSUSA-owned, shared-readable by all councils.', 'ext'),
 ('Vendor / Drop-ship Partners','External','Okta (separate domain)','Adobe Commerce (not AEM)',
  'Upload product images/files for their own products via Adobe Commerce only.',
  'Restricted to their own vendor folder. No access to other DAM areas.',
  'Minimal (metadata on own product assets).',
  '2-3 trusted partners','Dedicated vendor folder (like council structure). Daria reviews offline before approved assets are uploaded; no AEM review workflow.', 'ext'),
 ('Agency Partners','External','Okta (separate domain)','Content Hub / assigned',
  'Download select assigned folders only.',
  'Read + download on selected assigned folders only.',
  'None.',
  'Few','Download-only of assigned folders.', 'ext'),
 ('Creative Partners','External','Okta (separate domain)','Staging folder / Creative Cloud integration',
  'Upload finished assets to a staging folder only.',
  'Restricted to their own staging folder (upload + read).',
  'Minimal (metadata on own uploads).',
  '2-3','Assets created outside and uploaded. If using Adobe CC (Photoshop), AEM Assets integration available. Approved assets only.', 'ext'),
]
r = hrow+1
for vals in rows1:
    *cells, kind = vals
    put_row(ws1, r, cells, fill=(INT_FILL if kind=='int' else EXT_FILL), center_from=2)
    r += 1

widths1 = [22,14,16,22,40,34,26,14,46]
for i,w in enumerate(widths1, start=1): ws1.column_dimensions[get_column_letter(i)].width = w
ws1.freeze_panes = 'B5'
ws1.auto_filter.ref = f'A{hrow}:{get_column_letter(len(H1))}{hrow}'

# =====================================================================
# SHEET 2 - AEM Group / ACL Matrix (technical build grid)
# =====================================================================
ws2 = wb.create_sheet('AEM Group ACL Matrix')
ws2.cell(1,1,'GSUSA AEM Assets - Group / ACL Build Matrix').font = Font(bold=True, size=13, color='2C3E50')
ws2.cell(2,1,'R=Read, M=Modify, C=Create, D=Delete, R-ACL=Read ACL, E-ACL=Edit ACL, REPL=Replicate/Publish').font = Font(italic=True, size=9, color='566573')

H2 = ['AEM Group','Maps to User Type','Path (jcr)','Read','Modify','Create','Delete','Read ACL','Edit ACL','Replicate','Notes']
hrow2 = 4
for ci,h in enumerate(H2, start=1): ws2.cell(hrow2,ci,h)
style_header(ws2, hrow2, len(H2))

def aclrow(group, user, path, R='',M='',C='',D='',RA='',EA='',REPL='',notes=''):
    return [group,user,path,R,M,C,D,RA,EA,REPL,notes]

acl = [
 aclrow('gsusa-content-creators','GSUSA Content Creators','/content/dam/gsusa','R','M','C','D','','','REPL','Full superset on all folders'),
 aclrow('gsusa-content-managers','GSUSA Content Managers','/content/dam/gsusa/marketing','R','M','C','D','','','REPL','Write only on marketing'),
 aclrow('gsusa-content-managers','GSUSA Content Managers','/content/dam/gsusa','R','','','','','','','View-only on all other folders'),
 aclrow('gsusa-content-users','GSUSA Content Users','/content/dam/gsusa','R','','','','','','','Read + download only'),
 aclrow('gsusa-dam-admins','Admin / IT','/','R','M','C','D','R-ACL','E-ACL','REPL','Full admin incl. ACL + replicate'),
 aclrow('gsusa-tag-admins','Admin / IT (1-2)','/content/cq:tags/girlscouts','R','M','C','D','','','','Tag creation restricted to 1-2'),
 aclrow('council-{code}-collaborators','Council Collaborators','/content/dam/gsusa/councils/{code}','R','M','C','D','','','','Full within OWN council folder (one group per council)'),
 aclrow('council-{code}-collaborators','Council Collaborators','/content/dam/gsusa/global','R','','','','','','','View shared GSUSA global assets'),
 aclrow('council-{code}-collaborators','Council Collaborators','/content/dam/gsusa/products','R','','','','','','','View shared product assets'),
 aclrow('council-{code}-collaborators','Council Collaborators','/content/dam/gsusa/marketing','R','','','','','','','View shared marketing assets'),
 aclrow('council-{code}-collaborators','Council Collaborators','/content/dam/gsusa/councils/100','R','','','','','','','Girl Scout Central (GSUSA-owned, shared read)'),
 aclrow('vendor-{name}','Vendor / Drop-ship','/content/dam/gsusa/vendors/{name}','R','M','C','','','','','Own vendor folder only; primary path is Commerce'),
 aclrow('agency-{name}','Agency Partners','/content/dam/gsusa/{assigned}','R','','','','','','','Download-only, assigned folders'),
 aclrow('creative-{name}','Creative Partners','/content/dam/gsusa/staging/{name}','R','M','C','','','','','Upload to own staging folder only'),
]
r = hrow2+1
for vals in acl:
    put_row(ws2, r, vals, center_from=4)
    r += 1

widths2 = [28,22,34,7,8,8,8,9,9,10,44]
for i,w in enumerate(widths2, start=1): ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = 'D5'
ws2.auto_filter.ref = f'A{hrow2}:{get_column_letter(len(H2))}{hrow2}'

# =====================================================================
# SHEET 3 - Open Items / Assumptions
# =====================================================================
ws3 = wb.create_sheet('Open Items & Assumptions')
ws3.cell(1,1,'Open Items, Assumptions & Decisions to Confirm').font = Font(bold=True, size=13, color='2C3E50')
H3 = ['#','Item','Status / Assumption','Source']
hrow3 = 3
for ci,h in enumerate(H3, start=1): ws3.cell(hrow3,ci,h)
style_header(ws3, hrow3, len(H3))
items = [
 ('1','All external partners authenticated via Okta on separate (non-girlscouts.org) domains.','Confirmed in MOM 2.5','MOM'),
 ('2','Council Collaborators have NO direct AEM access; view/download (and possibly upload) via Content Hub.','Adobe to confirm CH upload capability','MOM / Workbook Q'),
 ('3','New tag creation restricted to 1-2 admins; all other roles apply existing tags only.','Confirmed','MOM 2.5'),
 ('4','One AEM group per council (council-{code}-collaborators); councils cannot see each others folders.','Confirmed','MOM 2.5'),
 ('5','Council 100 (Girl Scout Central) treated as GSUSA-owned, readable by all councils.','Confirmed','MOM 2.2/2.5'),
 ('6','Vendors upload via Adobe Commerce, not direct AEM; Daria reviews offline before upload.','Confirmed','MOM 2.5'),
 ('7','~20 internal users with full upload/edit access.','To confirm vs license terms','MOM 2.5'),
 ('8','Council users per council: at least 1, likely 2-3.','To confirm from prior docs','MOM 2.5'),
 ('9','Specific GSUSA folders to be hidden from council view.','Folders not yet identified','MOM 2.5'),
 ('10','Content Manager write scope limited to /marketing; view-only elsewhere.','Confirmed','Workbook + MOM'),
]
r = hrow3+1
for vals in items:
    put_row(ws3, r, vals, center_from=1)
    r += 1
for i,w in enumerate([5,52,40,16], start=1): ws3.column_dimensions[get_column_letter(i)].width = w
ws3.freeze_panes = 'A4'

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'GSUSA-Permission-Matrix.xlsx')
wb.save(out)
print('Saved:', out)
